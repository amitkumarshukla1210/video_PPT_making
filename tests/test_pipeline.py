"""End-to-end smoke test: screenshots + notes -> PPTX + script -> voice-over MP4 + SRT.
Run:  python -m pytest -q tests     (or)   python tests/test_pipeline.py"""
import json
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PIL import Image, ImageDraw  # noqa: E402
from pptx import Presentation     # noqa: E402

from app import project as P, tts                       # noqa: E402
from app.ffmpeg_utils import ffmpeg_bin                 # noqa: E402
from app.ppt_builder import build_pptx                  # noqa: E402
from app.ppt_import import import_pptx                  # noqa: E402
from app.script_writer import draft_all                 # noqa: E402
from app.video import render_video, estimate            # noqa: E402


def _fake_screenshot(path, label, color):
    im = Image.new("RGB", (1600, 900), "#f7f9fc")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1600, 70), fill=color)
    d.text((30, 25), label, fill="white")
    for i in range(6):
        d.rectangle((60 + i * 250, 600 - i * 70, 240 + i * 250, 820), fill="#1f5fa8")
    im.save(path)


def _duration(path):
    out = subprocess.run([ffmpeg_bin(), "-i", str(path)], capture_output=True, text=True).stderr
    hms = out.split("Duration: ")[1].split(",")[0]
    h, m, s = hms.split(":")
    return int(h) * 3600 + int(m) * 60 + float(s), out


def test_end_to_end():
    engine = tts.default_engine()
    assert engine, "no TTS engine installed"
    proj = P.new_project("Test Design to Cost")
    try:
        base = P.folder(proj["id"])
        proj.update(presenter="Amit Shukla", company="OEM Co.")
        proj["voice"] = {"engine": engine, "voice": "", "lang": "en-us", "speed": 1.0}
        for i, (name, notes) in enumerate([
            ("01_commodity-impact.png", "Steel price up 8% this quarter\nForex impact on imported parts\nBOM cost vs target"),
            ("02_part-costing.png", "Should-cost breakdown per part. Material, process and overhead."),
        ]):
            rel = f"uploads/{name}"
            _fake_screenshot(base / rel, name, ["#152238", "#e07b1f"][i])
            P.add_content_slide(proj, P.title_from_filename(name), notes, rel)
        P.add_content_slide(proj, "Text only slide", "No screenshot here\nJust bullets")
        draft_all(proj)
        assert all(s["script"] for s in proj["slides"])
        P.save(proj)

        pptx = build_pptx(proj)
        prs = Presentation(pptx)
        assert len(prs.slides) == 5
        assert "Steel price" in prs.slides[1].notes_slide.notes_text_frame.text

        # importing that deck back should yield 5 more slides with screenshots and scripts
        n = import_pptx(proj, pptx)
        assert n == 5
        proj["slides"] = proj["slides"][:5] + [proj["slides"][-1]]   # keep the test short

        est = estimate(proj)
        assert est["render_seconds"] > 0

        stages = []
        r = render_video(proj, progress=lambda st, f, m: stages.append(st))
        dur, info = _duration(r["video"])
        assert abs(dur - r["seconds"]) < 0.5, (dur, r["seconds"])
        assert "1920x1080" in info and "Audio: aac" in info and "mov_text" in info
        srt = Path(r["srt"]).read_text(encoding="utf-8")
        assert "-->" in srt and "Steel price" in srt
        print(json.dumps({"engine": engine, "video_s": round(dur, 1), "render_s": round(r["render_seconds"], 1),
                          "estimate": est}, indent=1))
    finally:
        P.delete(proj["id"])


if __name__ == "__main__":
    test_end_to_end()
    print("OK")
