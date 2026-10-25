"""Import an existing .pptx: each slide's title, text, largest picture and speaker notes
become a slide in the project, which is then rebuilt in the Design to Cost brand."""
import uuid

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .project import add_content_slide, folder


def _pictures(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield sh
        elif sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(sh.shapes)


def import_pptx(proj, pptx_path):
    base = folder(proj["id"])
    prs = Presentation(pptx_path)
    added = 0
    for n, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        texts.append(t)
        image = None
        pics = [p for p in _pictures(slide.shapes) if p.image.ext.lower() in ("png", "jpg", "jpeg", "gif", "bmp")]
        pics.sort(key=lambda p: (p.width or 0) * (p.height or 0), reverse=True)
        if pics:
            blob = pics[0].image
            name = f"uploads/import_{n:02d}_{uuid.uuid4().hex[:4]}.{blob.ext}"
            (base / name).write_bytes(blob.blob)
            image = name
        if not title:
            title = texts.pop(0) if texts else f"Slide {n}"
        s = add_content_slide(proj, title, "\n".join(texts[:8]), image)
        if slide.has_notes_slide:
            s["script"] = slide.notes_slide.notes_text_frame.text.strip()
        added += 1
    return added
