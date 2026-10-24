"""Build the branded .pptx (python-pptx, MIT licence)."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from . import layout as L
from .config import NAVY, BLUE, ORANGE, WHITE, LIGHT, GREY, FONT_NAME, DEFAULT_LOGO
from .project import folder


def _rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#"))


def _rect(slide, box, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, *(Inches(v) for v in box))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    if line:
        s.line.color.rgb = _rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, box, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(*(Inches(v) for v in box))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = FONT_NAME, Pt(size), bold
    r.font.color.rgb = _rgb(color)
    return tb


def _picture(slide, path, box, border=None):
    with Image.open(path) as im:
        w, h = im.size
    x, y, fw, fh = L.fit_box(w, h, box)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(fw), Inches(fh))
    if border:
        pic.line.color.rgb = _rgb(border)
        pic.line.width = Pt(1)
    return pic


def _logo(slide, logo_path, box):
    """White card + logo. Name the picture so it is easy to find and 'Change Picture' in PowerPoint."""
    card = _rect(slide, box, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    card.adjustments[0] = 0.18
    x, y, w, h = box
    pad = 0.08
    pic = _picture(slide, logo_path, (x + pad, y + pad, w - 2 * pad, h - 2 * pad))
    pic.name = "LOGO (right-click > Change Picture)"


def _bullets(slide, box, lines, size, fill=LIGHT, heading="Key points"):
    _rect(slide, box, fill)
    x, y, w, h = box
    _rect(slide, (x, y, 0.07, h), BLUE)
    tb = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.2), Inches(w - 0.4), Inches(h - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = heading
    r.font.name, r.font.size, r.font.bold = FONT_NAME, Pt(size + 2), True
    r.font.color.rgb = _rgb(BLUE)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(size * 0.6)
        r = p.add_run()
        r.text = "■  "
        r.font.name, r.font.size = FONT_NAME, Pt(size * 0.6)
        r.font.color.rgb = _rgb(ORANGE)
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = FONT_NAME, Pt(size)
        r.font.color.rgb = _rgb(NAVY)


def _footer(slide, proj, number):
    left = proj.get("title", "")
    if proj.get("company"):
        left += f"  |  {proj['company']}"
    _text(slide, (0.5, L.FOOTER_Y, 9, 0.35), left, 10, GREY)
    _text(slide, (11.8, L.FOOTER_Y, 1.03, 0.35), str(number), 10, BLUE, bold=True, align=PP_ALIGN.RIGHT)


def _cover(slide, proj, s, logo):
    _rect(slide, (0, 0, L.SLIDE_W, L.SLIDE_H), NAVY)
    _rect(slide, L.COVER_BAND, BLUE)
    _rect(slide, L.COVER_BAR, ORANGE)
    _logo(slide, logo, L.COVER_LOGO)
    _text(slide, L.COVER_TITLE, s["title"] or proj["title"], 44, WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _text(slide, L.COVER_SUB, s["notes"] or "", 22, "#bcd3f0")
    meta = "  |  ".join(v for v in (proj.get("presenter"), proj.get("company"), proj.get("contact")) if v)
    if meta:
        _text(slide, L.COVER_META, meta, 16, WHITE)


def _content(slide, proj, s, logo, base):
    _rect(slide, (0, 0, L.SLIDE_W, L.SLIDE_H), WHITE)
    _rect(slide, L.HEADER, NAVY)
    _rect(slide, L.ACCENT, ORANGE)
    _text(slide, L.HEADER_TITLE, s["title"], 28, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _logo(slide, logo, L.HEADER_LOGO)
    lines = L.bullets(s["notes"])
    img = base / s["image"] if s.get("image") else None
    if img and img.exists():
        _picture(slide, img, L.SHOT_WITH_NOTES if lines else L.SHOT_FULL, border="#c9d3e0")
        if lines:
            _bullets(slide, L.NOTES_PANEL, lines, L.bullet_font_pt(lines))
    elif lines:
        _bullets(slide, L.NOTES_FULL, lines, L.bullet_font_pt(lines, full_width=True))


def logo_path(proj):
    if proj.get("logo"):
        p = folder(proj["id"]) / proj["logo"]
        if p.exists():
            return p
    return DEFAULT_LOGO


def build_pptx(proj, out_path=None):
    base = folder(proj["id"])
    out_path = Path(out_path or base / "out" / "presentation.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(L.SLIDE_W), Inches(L.SLIDE_H)
    blank = prs.slide_layouts[6]
    logo = logo_path(proj)
    for i, s in enumerate(proj["slides"], start=1):
        slide = prs.slides.add_slide(blank)
        if s["kind"] in ("title", "closing"):
            _cover(slide, proj, s, logo)
        else:
            _content(slide, proj, s, logo, base)
            _footer(slide, proj, i)
        # The editable narration goes into the speaker notes.
        slide.notes_slide.notes_text_frame.text = s.get("script") or ""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
